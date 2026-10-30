"""Daily Excel ledger files (PRD FR-22 to FR-31; Architecture 5.1).

Each day with entries has its own file per user: <root>\\<YYYY>\\<MM-Mon>\\<DD-Mon-YYYY>_<username>.xlsx,
with one sheet named "Entries". Dates are stored as text (DD-Mon-YYYY), amounts as numbers.
Rows copied in from other users (CSV import) go to files named with their owner, beside the user's own.

Reading goes through a SQLite index of the files (entry_index.py): only files that changed on disk are opened again.

Saves go to a temporary file first and then replace the real file (Rules 16). If Excel
has the file open, Windows refuses the replace, and ExcelFileLockedError is raised so the
form can keep the entry (PRD FR-29, Rules 15).
"""

import os
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Callable

from app.core import dates
from app.core.models import DEFAULT_MAIN_ENTRIES, Entry
from app.storage.entry_index import EntryIndex

SHEET_NAME = "Entries"
HEADERS = ("ID", "Main Head", "Sub Head", "Date", "Name", "Particulars", "Debit", "Credit", "Created At")
ID_PATTERN = re.compile(r"^E(\d{8})-(\d{3})$")
CENT = Decimal("0.01")
_UNSAFE_FILE_CHARACTERS = re.compile(r'[^\w.\- ]')


def safe_owner(username: str) -> str:
    """The username as it appears in a file name: characters Windows does not allow become underscores."""
    return _UNSAFE_FILE_CHARACTERS.sub("_", username.strip()).strip(" .")


class ExcelFileLockedError(Exception):
    """The daily file is open in Excel, so it cannot be replaced. Nothing was lost."""


class _DefaultHeads:
    """Main Entry names for a store made without the Main Entry service (tests, tools)."""

    def label(self, main_head: str) -> str:
        return next((item["name"] for item in DEFAULT_MAIN_ENTRIES if item["id"] == main_head), main_head)

    def id_for_label(self, label: str) -> str:
        return next((item["id"] for item in DEFAULT_MAIN_ENTRIES if item["name"] == label), label)


class ExcelStore:
    """Reads and writes the daily workbook files under one root folder.

    heads turns a Main Entry id into the name saved in Excel and back (label() and id_for_label()).
    """

    def __init__(self, root: Path, heads=None, owner: Callable[[], str] | None = None, index_file: Path | None = None) -> None:
        self.root = root
        self.heads = heads or _DefaultHeads()
        self._owner = owner or (lambda: "")  # the logged-in username, asked each time because login comes after start-up
        # Parsed rows per file, reused until the file changes on disk (Architecture 7). index_file None: in memory only.
        self._index = EntryIndex(index_file, root)

    # ----- file locations -----

    def current_owner(self) -> str:
        return safe_owner(self._owner())

    def is_mine(self, entry: Entry) -> bool:
        """True for rows in the logged-in user's own files. Only those can be edited or deleted."""
        return entry.owner.lower() == self.current_owner().lower()

    def file_path_for(self, entry_date: date, owner: str | None = None) -> Path:
        owner = self.current_owner() if owner is None else safe_owner(owner)
        suffix = f"_{owner}" if owner else ""
        return self.root / dates.year_folder_name(entry_date) / dates.month_folder_name(entry_date) / f"{dates.format_date(entry_date)}{suffix}.xlsx"

    # ----- writing -----

    def next_id(self, entry_date: date) -> str:
        """The next ID for the day: E<YYYYMMDD>-<NNN>, one more than the highest number used that day."""
        highest = 0
        path = self.file_path_for(entry_date)
        if path.exists():
            _workbook, sheet = self._open(path)
            for row in sheet.iter_rows(min_row=2, values_only=True):
                match = ID_PATTERN.match(str(row[0] or ""))
                if match:
                    highest = max(highest, int(match.group(2)))
        return f"E{entry_date:%Y%m%d}-{highest + 1:03d}"

    def append_entry(self, entry: Entry) -> None:
        path = self.file_path_for(entry.entry_date, entry.owner or None)
        workbook, sheet = self._open(path) if path.exists() else self._new_workbook()
        self._append_row(sheet, self._to_row(entry))
        self._save(workbook, path)

    def import_entries(self, entries: list[Entry]) -> tuple[int, int]:
        """Copy rows from other users' CSV into files named with their owner. Returns (added, already_there).

        A row whose owner and ID are already in that day's file is skipped, so importing the same CSV twice adds nothing.
        """
        groups: dict[Path, list[Entry]] = {}
        for entry in entries:
            groups.setdefault(self.file_path_for(entry.entry_date, entry.owner or None), []).append(entry)
        added = skipped = 0
        for path, group in groups.items():
            workbook, sheet = self._open(path) if path.exists() else self._new_workbook()
            known = {str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None}
            fresh = [entry for entry in group if entry.id not in known]
            skipped += len(group) - len(fresh)
            for entry in fresh:
                self._append_row(sheet, self._to_row(entry))
            if fresh:
                self._save(workbook, path)
                added += len(fresh)
        return added, skipped

    def update_entry(self, entry: Entry, old_date: date) -> None:
        """Change a saved entry. If its date changed, the row moves to the file for the new date (FR-30)."""
        if old_date != entry.entry_date:
            self.delete_entry(entry.id, old_date)
            self.append_entry(entry)
            return
        path = self.file_path_for(old_date)
        workbook, sheet = self._open(path)
        row_number = self._find_row(sheet, entry.id)
        for column, value in enumerate(self._to_row(entry), start=1):
            self._set_cell(sheet.cell(row=row_number, column=column), value)
        self._save(workbook, path)

    def delete_entry(self, entry_id: str, entry_date: date) -> None:
        path = self.file_path_for(entry_date)
        workbook, sheet = self._open(path)
        sheet.delete_rows(self._find_row(sheet, entry_id))
        self._save(workbook, path)

    # ----- reading -----

    def read_all_entries(self) -> list[Entry]:
        """Every entry in every daily file under the root, oldest first. Only changed files are opened."""
        stamps: dict[str, tuple[int, int]] = {}
        if self.root.exists():
            for path in self.root.rglob("*.xlsx"):
                if ".tmp" not in path.name:
                    info = path.stat()
                    stamps[path.relative_to(self.root).as_posix()] = (info.st_mtime_ns, info.st_size)
        for rel in self._index.changed_since(stamps):
            path = self.root / rel
            _workbook, sheet = self._open(path)
            self._index.replace_file(rel, stamps[rel], self._parse_sheet(path, sheet))
        entries = self._index.all_entries()
        resolved: dict[str, str] = {}
        for entry in entries:  # the index holds the name written in Excel; turn it into the Main Entry id
            if entry.main_head not in resolved:
                resolved[entry.main_head] = self.heads.id_for_label(entry.main_head)
            entry.main_head = resolved[entry.main_head]
        return entries

    def _parse_sheet(self, path: Path, sheet) -> list[Entry]:
        owner = path.stem.split("_", 1)[1] if "_" in path.stem else ""  # the date part has no underscore
        return [self._from_row(row, owner) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None]

    # ----- helpers -----

    @staticmethod
    def _new_workbook():
        import openpyxl  # loaded when first needed: a start that reads only the index never opens a workbook
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(list(HEADERS))
        return workbook, sheet

    @classmethod
    def _append_row(cls, sheet, values: list) -> None:
        row_number = sheet.max_row + 1
        for column, value in enumerate(values, start=1):
            cls._set_cell(sheet.cell(row=row_number, column=column), value)

    @staticmethod
    def _set_cell(cell, value) -> None:
        """Write a value. Text starting with = stays text; Excel would otherwise run it as a formula."""
        cell.value = value
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"

    @staticmethod
    def _open(path: Path):
        import openpyxl
        workbook = openpyxl.load_workbook(path)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
        return workbook, sheet

    @staticmethod
    def _find_row(sheet, entry_id: str) -> int:
        for row_number in range(2, sheet.max_row + 1):
            if sheet.cell(row=row_number, column=1).value == entry_id:
                return row_number
        raise KeyError(f"entry {entry_id} not found")

    def _save(self, workbook, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = path.with_name(path.stem + ".tmp.xlsx")
        workbook.save(temp_path)
        try:
            os.replace(temp_path, path)
        except PermissionError as error:
            temp_path.unlink(missing_ok=True)
            raise ExcelFileLockedError(str(error)) from error
        self._reindex(path, workbook)

    def _reindex(self, path: Path, workbook) -> None:
        """Update the index from the workbook just saved, so the next read does not open the file again."""
        rel = path.relative_to(self.root).as_posix()
        try:
            sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
            info = path.stat()
            self._index.replace_file(rel, (info.st_mtime_ns, info.st_size), self._parse_sheet(path, sheet))
        except Exception:  # the file is saved; at worst the next read opens it again
            self._index.forget_file(rel)

    def _to_row(self, entry: Entry) -> list:
        created = f"{dates.format_date(entry.created_at.date())} {entry.created_at:%H:%M}"
        return [
            entry.id,
            self.heads.label(entry.main_head),
            entry.sub_head,
            dates.format_date(entry.entry_date),
            entry.name,
            entry.particulars,
            float(entry.debit),
            float(entry.credit),
            created,
        ]

    def _from_row(self, row, owner: str = "") -> Entry:
        """One row as an Entry. main_head is still the name written in Excel; read_all_entries turns it into the id."""
        main_head = str(row[1])
        created_text = str(row[8])
        created_date_text, created_time_text = created_text.rsplit(" ", 1)
        created_date = dates.parse_date(created_date_text)
        hour, minute = (int(part) for part in created_time_text.split(":"))
        return Entry(
            id=str(row[0]),
            main_head=main_head,
            sub_head=str(row[2]),
            entry_date=dates.parse_date(str(row[3])),
            name=str(row[4]),
            particulars=str(row[5]),
            debit=Decimal(str(row[6])).quantize(CENT),
            credit=Decimal(str(row[7])).quantize(CENT),
            created_at=datetime(created_date.year, created_date.month, created_date.day, hour, minute),
            owner=owner,
        )
