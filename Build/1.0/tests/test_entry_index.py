"""Tests for the SQLite index of the Excel files (app/storage/entry_index.py)."""

import os
import time
from datetime import date
from decimal import Decimal

import openpyxl

from app.services.entry_service import EntryService
from app.storage.excel_store import ExcelStore


def make_store(root, index_file):
    return ExcelStore(root, owner=lambda: "asha", index_file=index_file)


def save(store, name="Lenin", entry_date=date(2026, 10, 20)):
    EntryService(store).save_entry("MH 2", "Lenin", entry_date, name, "Morning", Decimal("10.00"), Decimal("0.00"))


def count_opens(monkeypatch):
    opened = []
    real = openpyxl.load_workbook
    monkeypatch.setattr(openpyxl, "load_workbook", lambda path, *a, **k: (opened.append(path), real(path, *a, **k))[1])
    return opened


def test_a_new_session_reads_from_the_index_without_opening_workbooks(tmp_path, monkeypatch):
    root, index = tmp_path / "data", tmp_path / "index.db"
    first = make_store(root, index)
    save(first)
    save(first, "Vivek", date(2026, 10, 21))
    first._index.close()

    opened = count_opens(monkeypatch)
    second = make_store(root, index)
    assert [entry.name for entry in second.read_all_entries()] == ["Lenin", "Vivek"]
    assert opened == []


def test_saving_updates_the_index_so_the_next_read_opens_nothing(tmp_path, monkeypatch):
    store = make_store(tmp_path / "data", None)
    save(store)
    opened = count_opens(monkeypatch)
    assert len(store.read_all_entries()) == 1
    save(store, "Second")
    opened.clear()
    assert len(store.read_all_entries()) == 2
    assert opened == []


def test_a_file_edited_outside_the_app_is_read_again(tmp_path):
    store = make_store(tmp_path / "data", tmp_path / "index.db")
    save(store)
    assert store.read_all_entries()[0].name == "Lenin"
    path = store.file_path_for(date(2026, 10, 20))
    workbook = openpyxl.load_workbook(path)
    workbook["Entries"]["E2"] = "Edited in Excel"
    workbook.save(path)
    os.utime(path, ns=(time.time_ns() + 5_000_000_000,) * 2)
    assert store.read_all_entries()[0].name == "Edited in Excel"


def test_a_removed_file_disappears_from_the_list(tmp_path):
    store = make_store(tmp_path / "data", tmp_path / "index.db")
    save(store)
    save(store, "Vivek", date(2026, 10, 21))
    store.file_path_for(date(2026, 10, 21)).unlink()
    assert [entry.name for entry in store.read_all_entries()] == ["Lenin"]


def test_a_damaged_index_is_rebuilt_from_the_excel_files(tmp_path):
    root, index = tmp_path / "data", tmp_path / "index.db"
    first = make_store(root, index)
    save(first)
    first._index.close()
    index.write_bytes(b"this is not a database" * 50)
    assert [entry.name for entry in make_store(root, index).read_all_entries()] == ["Lenin"]


def test_a_different_data_folder_does_not_reuse_the_index(tmp_path):
    index = tmp_path / "index.db"
    first = make_store(tmp_path / "one", index)
    save(first)
    first._index.close()
    assert make_store(tmp_path / "two", index).read_all_entries() == []
